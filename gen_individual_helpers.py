"""
Generate the 3 Individual R3 Reports based on the updated Review3_Unified_Master_Report.docx
and the individual_report_editable format, with 10 role-specific Q1 literature survey papers each,
exact table styling (shd fill F2F2F2 header/labels, F2F2F2 metadata, F0F4F8 callouts, white data cells).
"""
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    if fill_hex:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, border_color="D3D3D3"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:left w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/><w:right w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/></w:tcBorders>')
    tcPr.append(borders)

def format_doc_table(table, col_widths=None, is_metadata=False):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    if is_metadata:
        for row in table.rows:
            for cell in row.cells:
                set_cell_shading(cell, "F2F2F2")
                set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
                set_cell_borders(cell, "D3D3D3")
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
    else:
        # Header row: F2F2F2
        for cell in table.rows[0].cells:
            set_cell_shading(cell, "F2F2F2")
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            set_cell_borders(cell, "D3D3D3")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.bold = True
        # Data rows: FFFFFF
        for row in table.rows[1:]:
            for cell in row.cells:
                set_cell_shading(cell, "FFFFFF")
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                set_cell_borders(cell, "D3D3D3")
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width

def add_callout(doc, title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.rows[0].cells[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, "F0F4F8")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="0" w:color="1B365D"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_body = p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Setup helper functions ready")
