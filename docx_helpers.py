"""
Generate comprehensive Review 3 Unified Document (docx) integrating R1, R2, R3 rubrics.
"""
import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shading if any
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = ['<w:tcBorders ' + nsdecls("w") + '>']
    for side, border in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if border:
            borders.append(f'<w:{side} w:val="{border.get("val", "single")}" w:sz="{border.get("sz", "4")}" w:space="0" w:color="{border.get("color", "D3D3D3")}"/>')
        else:
            borders.append(f'<w:{side} w:val="none"/>')
    borders.append('</w:tcBorders>')
    tcPr.append(parse_xml(''.join(borders)))

def format_table(table, col_widths=None, header_bg="1B365D", alt_bg="F9FBFD"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header row
    for col_idx, cell in enumerate(table.rows[0].cells):
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        border = {"val": "single", "sz": "6", "color": "1B365D"}
        set_cell_borders(cell, top=border, bottom=border, left=border, right=border)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                
    # Data rows
    for row_idx, row in enumerate(table.rows[1:], start=1):
        bg = alt_bg if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell in enumerate(row.cells):
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            border = {"val": "single", "sz": "4", "color": "D3D3D3"}
            set_cell_borders(cell, top=border, bottom=border, left=border, right=border)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(51, 51, 51)
                    
    # Column widths
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width

def add_callout_box(doc, text, title="KEY HIGHLIGHT", bg_color="F0F4F8", border_color="1B365D"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.rows[0].cells[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"[{title}] ")
    run_t.font.name = "Calibri"
    run_t.font.bold = True
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(27, 54, 93) # Navy
    
    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(51, 51, 51)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Setup helper functions complete")
